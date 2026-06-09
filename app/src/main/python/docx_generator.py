"""
Generator obdukcionog zapisnika iz template-a.

VAŽNO: Template fajl se NIKAD ne mijenja. Učitava se u memoriju,
modificira se kopija, snima se kao novi fajl u OUTPUT_DIR.
"""
from pathlib import Path
from datetime import datetime
import re
import copy

from docx import Document
from docx.oxml.ns import qn
from copy import deepcopy

from sections import DICTATION_SECTIONS, HEADER_FIELDS, DEFAULT_IZUZETI_UZORCI


def _ensure_run_format_from_pPr(para):
    """Ako paragraf ima pPr->rPr (npr. Georgia font na nivou paragrafa) ali
    runs nemaju vlastiti rPr, Word ignoriše pPr font i pada na style default
    (Calibri 11). Kopiramo pPr->rPr na svaki run koji nema vlastiti rPr."""
    pPr = para._p.find(qn('w:pPr'))
    if pPr is None:
        return
    pPr_rPr = pPr.find(qn('w:rPr'))
    if pPr_rPr is None:
        return
    for r in para._p.findall(qn('w:r')):
        if r.find(qn('w:rPr')) is None:
            new_rPr = deepcopy(pPr_rPr)
            r.insert(0, new_rPr)  # rPr mora biti prvi child od <w:r>


USER_CONTENT_HIGHLIGHT = "yellow"  # Word highlight color (named values only)


def _mark_run_as_user_content(run_element, force_normal_weight=True):
    """Označi run kao korisnikov unos sa Word highlight bojom (žutom).
    Koristi <w:highlight> (Word standard highlight) umjesto <w:shd> — lakše ga je
    ukloniti u Word-u kroz "Highlight: None" toolbar opciju.

    Ako force_normal_weight=True, dodatno uklanja bold/underline (jer u sekcijama
    template ima bold samo na "5." prefiksu pa kompletni paragraf nasljeđuje bold).
    Ako False, čuva bold/underline kako je u template-u (za tabelu zaglavlja gdje
    vrijednost treba nasljediti bold cijele ćelije, npr. DŽINDO ĆAMIL).
    """
    from docx.oxml import OxmlElement
    rPr = run_element.find(qn('w:rPr'))
    if rPr is None:
        rPr = OxmlElement('w:rPr')
        run_element.insert(0, rPr)
    # Ukloni postojeći highlight i shading (ako su ostali iz prethodnih verzija)
    for tag in ('w:highlight', 'w:shd'):
        for el in rPr.findall(qn(tag)):
            rPr.remove(el)
    # Dodaj highlight = yellow (Word "Text Highlight Color" toolbar opcija)
    h = OxmlElement('w:highlight')
    h.set(qn('w:val'), USER_CONTENT_HIGHLIGHT)
    rPr.append(h)

    if force_normal_weight:
        # Ukloni bold/underline (za sekcije u tijelu)
        for tag in ('w:b', 'w:bCs'):
            for el in rPr.findall(qn(tag)):
                rPr.remove(el)
        b = OxmlElement('w:b'); b.set(qn('w:val'), '0'); rPr.append(b)
        bCs = OxmlElement('w:bCs'); bCs.set(qn('w:val'), '0'); rPr.append(bCs)
        for el in rPr.findall(qn('w:u')):
            rPr.remove(el)
        u = OxmlElement('w:u'); u.set(qn('w:val'), 'none'); rPr.append(u)


def _mark_paragraph_as_user_content(para, force_normal_weight=True):
    """Primijeni _mark_run_as_user_content na sve runs paragrafa."""
    for r in para._p.findall(qn('w:r')):
        _mark_run_as_user_content(r, force_normal_weight=force_normal_weight)


def _set_cell_label_value(cell, label: str, value: str):
    """Postavi prvi paragraf ćelije sa: label (template format, bez highlight)
    + value (sa svjetlom žutom pozadinom). Tako se jasno vidi šta je korisnik unio."""
    from docx.oxml import OxmlElement
    paragraphs = list(cell.paragraphs)
    if not paragraphs:
        cell.add_paragraph(label + ((" " + value) if value else ""))
        return
    para = paragraphs[0]
    runs = list(para._p.findall(qn('w:r')))

    if not runs:
        # Kreiraj novi run sa label-om
        para.add_run(label)
        runs = list(para._p.findall(qn('w:r')))

    first_r = runs[0]
    # Obriši ostale postojeće runs
    for r in runs[1:]:
        first_r.getparent().remove(r)
    # Postavi label u prvi run — NE dira format (bold ako je bio, ostaje)
    for t_el in first_r.findall(qn('w:t')):
        first_r.remove(t_el)
    t = OxmlElement('w:t')
    t.text = label
    if label != label.strip():
        t.set(qn('xml:space'), 'preserve')
    first_r.append(t)

    # Osiguraj da label run ima font (Georgia) iz pPr
    _ensure_run_format_from_pPr(para)

    if not value:
        return

    # Kloniraj prvi run da naslijedi font format, ALI bez highlight još
    new_r = deepcopy(first_r)
    # Obriši tekst iz klonirane verzije
    for t_el in new_r.findall(qn('w:t')):
        new_r.remove(t_el)
    # Ako kloniran run ima shading/highlight u rPr (od ranijeg generisanja), ukloni
    new_rPr = new_r.find(qn('w:rPr'))
    if new_rPr is not None:
        for tag in ('w:shd', 'w:highlight'):
            for el in new_rPr.findall(qn(tag)):
                new_rPr.remove(el)
    # Stavi vrijednost
    t = OxmlElement('w:t')
    t.text = " " + value
    t.set(qn('xml:space'), 'preserve')
    new_r.append(t)
    # Označi sa žutom pozadinom — NE diramo bold (vrijednost nasleđuje bold od ćelije,
    # npr. "DŽINDO ĆAMIL" ostaje bold ako je cijela ćelija bold u template-u)
    _mark_run_as_user_content(new_r, force_normal_weight=False)
    # Ubaci novi run odmah iza label run-a
    first_r.addnext(new_r)


def _replace_paragraph_text(para, new_text: str, mark_user=True):
    """Briše sve runs u paragrafu i ubacuje jedan novi run sa zadanim tekstom.
    Čuva stil paragrafa (Normal, List Paragraph, itd) i font.
    Ako mark_user=True, dodaje žuti highlight (označava korisnikov unos)."""
    runs = list(para.runs)
    if runs:
        first_run = runs[0]
        for r in runs[1:]:
            r._r.getparent().remove(r._r)
        first_run.text = new_text
    else:
        para.add_run(new_text)
    _ensure_run_format_from_pPr(para)
    if mark_user:
        _mark_paragraph_as_user_content(para)


def _insert_paragraph_after(para, text: str, style=None):
    """Ubaci novi paragraf odmah ispod 'para'. Čuva sve formatiranje
    (font, veličinu, alignment) klonirajući XML i zadržavajući prvi run."""
    from docx.oxml import OxmlElement
    new_p_xml = deepcopy(para._p)
    runs = new_p_xml.findall(qn('w:r'))
    if runs:
        # Zadrži PRVI run sa svim format-om (rPr), obriši ostale
        for r in runs[1:]:
            new_p_xml.remove(r)
        first_r = runs[0]
        # Obriši postojeće tekst elemente, zadrži rPr (run properties = font, size, bold...)
        for t_el in first_r.findall(qn('w:t')):
            first_r.remove(t_el)
        for tab_el in first_r.findall(qn('w:tab')):
            first_r.remove(tab_el)
        for br_el in first_r.findall(qn('w:br')):
            first_r.remove(br_el)
        # Dodaj novi tekst — handle tab characters
        if "\t" in text:
            parts = text.split("\t")
            for i, part in enumerate(parts):
                if i > 0:
                    first_r.append(OxmlElement('w:tab'))
                if part:
                    t = OxmlElement('w:t')
                    t.text = part
                    if part != part.strip():
                        t.set(qn('xml:space'), 'preserve')
                    first_r.append(t)
        else:
            t = OxmlElement('w:t')
            t.text = text
            if text != text.strip():
                t.set(qn('xml:space'), 'preserve')
            first_r.append(t)
    else:
        # Nema runs u original-u — kreiraj novi
        r = OxmlElement('w:r')
        t = OxmlElement('w:t')
        t.text = text
        r.append(t)
        new_p_xml.append(r)

    para._p.addnext(new_p_xml)
    from docx.text.paragraph import Paragraph
    new_para = Paragraph(new_p_xml, para._parent)
    if style:
        new_para.style = style
    _ensure_run_format_from_pPr(new_para)
    _mark_paragraph_as_user_content(new_para)
    return new_para


def _clear_paragraph(para):
    for r in list(para.runs):
        r._r.getparent().remove(r._r)


def _fill_header_table(doc, header_data: dict):
    """Popunjava tabelu zaglavlja: dodaje vrijednosti uz postojeće labele."""
    if not doc.tables:
        return
    table = doc.tables[0]

    def set_cell(row_idx, col_idx, label, value):
        try:
            cell = table.rows[row_idx].cells[col_idx]
        except IndexError:
            return
        # Label = template format (ostaje kako je u template-u, bez highlight)
        # Value = svjetla žuta pozadina (vidi se da je korisnik unio)
        _set_cell_label_value(cell, label, value or "")

    ime = header_data.get("ime_prezime", "").strip()
    if ime:
        ime = ime.upper()

    set_cell(1, 0, "Ime i prezime:", ime)
    set_cell(1, 2, "Državljanin:", header_data.get("drzavljanin", ""))
    set_cell(2, 0, "Prebivalište:", header_data.get("prebivaliste", ""))
    set_cell(2, 2, "Adresa:", header_data.get("adresa", ""))
    set_cell(3, 0, "Rođen/a:", header_data.get("rodjen", ""))
    set_cell(3, 2, "Pronađen/preminuo:", header_data.get("pronadjen", ""))

    # Cell (4,1): "Kantonalnog Tužilaštva..." (statič) + "tužilac: NAME" + "veza: BROJ" + "Naredbom..." (statič)
    cell_41 = table.rows[4].cells[1]
    tuzilac = header_data.get("tuzilac", "")
    kt_broj = header_data.get("kt_broj", "")
    if kt_broj and not kt_broj.startswith("T09"):
        kt_broj_full = f"T09 0 KTA {kt_broj}".strip()
    else:
        kt_broj_full = kt_broj
    _set_cell_lines_smart(cell_41, [
        {"static": "Kantonalnog Tužilaštva Kantona Sarajevo"},
        {"label": "tužilac:", "value": tuzilac},
        {"label": "veza:", "value": kt_broj_full},
        {"static": "Naredbom tužioca naložena je obdukcija leša"},
    ])

    # Cell (5,0): label - može biti "Okolnosti slučaja:" ili "Uviđaj:"
    okolnosti_label = header_data.get("okolnosti_label", "Okolnosti slučaja").strip() or "Okolnosti slučaja"
    set_cell(5, 0, f"{okolnosti_label}:", "")

    # Cell (5,1): okolnosti (USER) → izuzeti uzorci (USER) → standardni dodatak (statič)
    cell_51 = table.rows[5].cells[1]
    okolnosti_tekst = header_data.get("okolnosti", "").strip()
    # Prazno (ništa nije izuzeto) → eksplicitna napomena, NE standardni spisak iz template-a
    izuzeti_tekst = (header_data.get("izuzeti_uzorci") or "").strip() or "Nisu izuzeti uzorci za dodatne pretrage."
    standardni_obdukcija = (
        "Obdukcija obavljena u prosekturi Katedre za sudsku medicinu Medicinskog Fakulteta UNSA. "
        "Pregled tijela fotografisao krim-tehničar."
    )
    specs_51 = []
    if okolnosti_tekst:
        for ln in okolnosti_tekst.split("\n"):
            if ln.strip():
                specs_51.append({"user": ln.strip()})
        specs_51.append({"static": ""})  # prazna linija razmaka
    if izuzeti_tekst:
        for ln in izuzeti_tekst.split("\n"):
            if ln.strip():
                specs_51.append({"user": ln.strip()})
        specs_51.append({"static": ""})
    specs_51.append({"static": standardni_obdukcija})
    _set_cell_lines_smart(cell_51, specs_51)

    # Cell (6,0): Obdukovan/a: datum
    set_cell(6, 0, "Obdukovan/a:", header_data.get("datum_obdukcije", ""))

    # Cell (6,2): Obducent (statič) + Pomoćnik (label/value)
    cell_62 = table.rows[6].cells[2]
    pomocnik = header_data.get("pomocnik", "")
    _set_cell_lines_smart(cell_62, [
        {"static": "Obducent: Prof. dr. Adis Salihbegović"},
        {"label": "Pomoćnik obducenta:", "value": pomocnik},
    ])


def _clone_paragraph_with_text(template_para, text: str):
    """Klonira XML postojećeg paragrafa (čuva pPr + rPr), zamijeni tekst.
    Vraća novi <w:p> XML element koji možeš ubaciti u dokument."""
    from docx.oxml import OxmlElement
    new_p = deepcopy(template_para._p)
    runs = new_p.findall(qn('w:r'))
    if runs:
        for r in runs[1:]:
            new_p.remove(r)
        first_r = runs[0]
        for t_el in first_r.findall(qn('w:t')):
            first_r.remove(t_el)
        t = OxmlElement('w:t')
        t.text = text
        if text != text.strip():
            t.set(qn('xml:space'), 'preserve')
        first_r.append(t)
    else:
        r = OxmlElement('w:r')
        t = OxmlElement('w:t')
        t.text = text
        r.append(t)
        new_p.append(r)
    # Osiguraj run-level rPr (kopiraj iz pPr->rPr ako fali)
    pPr = new_p.find(qn('w:pPr'))
    if pPr is not None:
        pPr_rPr = pPr.find(qn('w:rPr'))
        if pPr_rPr is not None:
            for r in new_p.findall(qn('w:r')):
                if r.find(qn('w:rPr')) is None:
                    r.insert(0, deepcopy(pPr_rPr))
    # NE mark-uj ovdje — pozivaoc odlučuje da li da mark-uje
    return new_p


def _set_cell_lines_smart(cell, specs):
    """Multiline cell sa specifikacijom šta je statič vs šta je korisnikov unos.
    specs je lista dictova, jedan po liniji:
      {"static": "tekst"}              — template format, bez highlight
      {"label": "...", "value": "..."} — label bez highlight, value sa highlight
      {"user": "tekst"}                — cijeli tekst je korisnikov, highlight
    """
    paragraphs = list(cell.paragraphs)
    if not paragraphs:
        # Bez paragrafa u ćeliji — kreiraj prvi
        cell.add_paragraph("")
        paragraphs = list(cell.paragraphs)

    # Prvi paragraf je referenca formata
    first = paragraphs[0]
    # Obriši ostale postojeće paragrafe
    for p in paragraphs[1:]:
        p._p.getparent().remove(p._p)

    if not specs:
        _replace_paragraph_text(first, "", mark_user=False)
        return

    def _strip_shd(para):
        """Ukloni shading i highlight sa svih runs paragrafa (nasleđen od kloniranog template-a)."""
        for r in para._p.findall(qn('w:r')):
            rPr = r.find(qn('w:rPr'))
            if rPr is not None:
                for tag in ('w:shd', 'w:highlight'):
                    for el in rPr.findall(qn(tag)):
                        rPr.remove(el)

    def apply_spec_to_paragraph(target_para, spec):
        if "static" in spec:
            _replace_paragraph_text(target_para, spec["static"], mark_user=False)
            _strip_shd(target_para)
        elif "user" in spec:
            # Multi-line user content u tabeli (okolnosti, izuzeti) — čuva bold ako je ćelija bold
            _replace_paragraph_text(target_para, spec["user"], mark_user=False)
            for r in target_para._p.findall(qn('w:r')):
                _mark_run_as_user_content(r, force_normal_weight=False)
        elif "label" in spec:
            label = spec.get("label", "")
            value = spec.get("value", "")
            _replace_paragraph_text(target_para, label, mark_user=False)
            _strip_shd(target_para)  # ukloni nasleđen shading sa label run-a
            if value:
                runs = target_para._p.findall(qn('w:r'))
                if runs:
                    first_r = runs[0]
                    new_r = deepcopy(first_r)
                    for t_el in new_r.findall(qn('w:t')):
                        new_r.remove(t_el)
                    new_rPr = new_r.find(qn('w:rPr'))
                    if new_rPr is not None:
                        for shd in new_rPr.findall(qn('w:shd')):
                            new_rPr.remove(shd)
                    from docx.oxml import OxmlElement
                    t = OxmlElement('w:t')
                    t.text = " " + value
                    t.set(qn('xml:space'), 'preserve')
                    new_r.append(t)
                    # Vrijednost nasljeđuje bold od ćelije template-a (force_normal_weight=False)
                    _mark_run_as_user_content(new_r, force_normal_weight=False)
                    first_r.addnext(new_r)

    # Prvi spec ide u postojeći prvi paragraf (zadržava format template-a)
    apply_spec_to_paragraph(first, specs[0])

    # Naredni spec-ovi ubacuju nove paragrafe ispod
    last = first
    for spec in specs[1:]:
        new_p_xml = _clone_paragraph_with_text(first, "")
        last._p.addnext(new_p_xml)
        from docx.text.paragraph import Paragraph
        new_para = Paragraph(new_p_xml, first._parent)
        apply_spec_to_paragraph(new_para, spec)
        last = new_para


def _set_cell_lines(cell, lines):
    """Resetuje sve paragrafe ćelije i popunjava sa zadatim listom linija.
    Cijeli sadržaj je tretiran kao korisnikov unos (svjetli highlight).
    KLJUČNO: čuva formatiranje (font, veličinu) iz template-a klonirajući XML."""
    paragraphs = list(cell.paragraphs)
    if not paragraphs:
        # Edge case — ćelija nema ni jedan paragraf (rijetko)
        cell.add_paragraph(lines[0] if lines else "")
        for ln in lines[1:]:
            cell.add_paragraph(ln)
        return

    first = paragraphs[0]
    # Prvi paragraf — koristi za template format. Zamijeni mu tekst.
    _replace_paragraph_text(first, lines[0] if lines else "")

    # Obriši ostale postojeće paragrafe (ali ne i prvi — on je referenca formatiranja)
    for p in paragraphs[1:]:
        p._p.getparent().remove(p._p)

    # Za svaku narednu liniju — kloniraj XML prvog paragrafa (koji nosi font format)
    # i ubaci sa novim tekstom. Tako svi paragrafi imaju isti font kao template.
    last_xml = first._p
    for ln in lines[1:]:
        new_p_xml = _clone_paragraph_with_text(first, ln)
        last_xml.addnext(new_p_xml)
        last_xml = new_p_xml


def _fill_section(para, section: dict, value: str):
    """Popunjava jednu diktiranu sekciju u tijelu dokumenta.

    Prima konkretni Paragraph objekt (stabilna referenca), ne indeks —
    zato što insertovanje paragrafa pomjera sve indekse nakon njega.
    """
    if para is None:
        return

    multi = section.get("multi", False)
    numbered = section.get("numbered", False)

    text = (value or "").strip()

    # Ekstraktiraj broj/prefix iz originalnog template teksta (npr. "1. ", "5. ", "11. ")
    original = para.text
    prefix_match = re.match(r"^(\s*\d+\.\s*)", original)
    number_prefix = prefix_match.group(1) if prefix_match else ""

    if not multi:
        # Jedan paragraf — prepiši
        new_text = f"{number_prefix}{text}" if number_prefix and not text.startswith(number_prefix.strip()) else text
        _replace_paragraph_text(para, new_text)
        return

    # Multi — više paragrafa
    lines = [ln.strip() for ln in text.split("\n") if ln.strip()]

    if numbered:
        lines = [f"{i+1}.\t{ln}" for i, ln in enumerate(lines)]
    else:
        # Prvi paragraf dobija prefix sekcije ako postoji
        if lines and number_prefix:
            lines[0] = f"{number_prefix}{lines[0]}"

    if not lines:
        _replace_paragraph_text(para, number_prefix.rstrip() if number_prefix else "")
        return

    # Prvi line — u postojeći paragraf
    _replace_paragraph_text(para, lines[0])

    # Ostatak — ubaci nove paragrafe ispod, redom
    last = para
    for ln in lines[1:]:
        last = _insert_paragraph_after(last, ln, style=para.style.name)


def generate_report(template_path: str, output_dir: str,
                    header_data: dict, sections_data: dict) -> Path:
    """Glavni entry point.

    template_path: putanja do .docx template-a (READ-ONLY, nikad se ne mijenja)
    output_dir: folder gdje se snima generisani fajl
    header_data: dict sa header poljima (ime_prezime, kt_broj, ...)
    sections_data: dict mapping section_id -> dictated text

    Vraća putanju generisanog fajla.
    """
    template_path = Path(template_path)
    if not template_path.exists():
        raise FileNotFoundError(f"Template ne postoji: {template_path}")

    # Učitaj template (python-docx učitava u memoriju, ne otvara fajl drveno)
    doc = Document(str(template_path))

    # KORAK 1: Sačuvaj reference na sve paragrafe koje ćemo modifikovati
    # PRIJE bilo kakvih izmjena. Insertovanje paragrafa pomjera indekse,
    # pa moramo raditi sa stabilnim XML referencama.
    section_paragraphs = {}
    for section in DICTATION_SECTIONS:
        idx = section["para_idx"]
        if idx < len(doc.paragraphs):
            section_paragraphs[section["id"]] = doc.paragraphs[idx]

    # KORAK 2: Popuni zaglavlje (ne dira tijelo dokumenta)
    _fill_header_table(doc, header_data)

    # KORAK 3: Popuni svaku sekciju koristeći sačuvane reference
    for section in DICTATION_SECTIONS:
        sid = section["id"]
        if sid in sections_data and sid in section_paragraphs:
            _fill_section(section_paragraphs[sid], section, sections_data[sid])

    # Generiši ime fajla
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    # Naming konvencija foldera predmeta: "Ime Prezime - T09 0 KTA broj.docx"
    ime = (header_data.get("ime_prezime") or "Bez imena").strip()
    kt = (header_data.get("kt_broj") or "").strip()
    if kt and not kt.upper().startswith("T"):
        kt_full = f"T09 0 KTA {kt}"
    else:
        kt_full = kt
    # Sanitiziraj ime za fajl (zadrži dijakritike i razmake, samo izbaci nevalidne windows znake)
    invalid = r'<>:"/\|?*'
    safe_ime = "".join(c for c in ime if c not in invalid).strip()
    safe_kt = "".join(c for c in kt_full if c not in invalid).strip()
    if safe_ime and safe_kt:
        base = f"{safe_ime} - {safe_kt}"
    elif safe_ime:
        base = safe_ime
    elif safe_kt:
        base = safe_kt
    else:
        # Fallback — koristi datum-vrijeme da bar ne bude "Zapisnik.docx" duplikat
        base = "Zapisnik " + datetime.now().strftime("%Y-%m-%d %H-%M")

    output_path = output_dir / f"{base}.docx"
    # Ako postoji, dodaj redni broj (- 1, - 2, - 3, ...) da ne prepiše prethodni
    if output_path.exists():
        n = 1
        while True:
            candidate = output_dir / f"{base} - {n}.docx"
            if not candidate.exists():
                output_path = candidate
                break
            n += 1
            if n > 999:  # safety
                output_path = output_dir / f"{base} - {datetime.now().strftime('%H%M%S')}.docx"
                break

    doc.save(str(output_path))
    return output_path
